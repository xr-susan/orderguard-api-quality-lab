"""Excel adapter for the canonical case model."""

from __future__ import annotations

import json
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ._common import case_id_from, validate_case_mapping
from .models import CaseDataError, CaseSpec

_STRUCTURED_FIELDS = {
    "request",
    "expected",
    "request.headers",
    "request.query",
    "request.params",
    "request.json",
    "request.body",
    "request.payload",
    "expected.headers",
    "expected.json",
    "expected.json_equals",
    "expected.json_contains",
    "expected.contains",
    "expected.json_paths",
    "expected.schema",
    "headers",
    "query",
    "params",
    "json",
    "body",
    "payload",
    "expected_headers",
    "expected_json",
    "json_equals",
    "json_contains",
    "json_paths",
    "schema",
}

_FLAT_TARGETS = {
    "method": ("request", "method"),
    "path": ("request", "path"),
    "headers": ("request", "headers"),
    "query": ("request", "query"),
    "params": ("request", "query"),
    "json": ("request", "json"),
    "body": ("request", "json"),
    "payload": ("request", "json"),
    "timeout_seconds": ("request", "timeout_seconds"),
    "status": ("expected", "status"),
    "status_code": ("expected", "status"),
    "expected_headers": ("expected", "headers"),
    "expected_json": ("expected", "json"),
    "json_equals": ("expected", "json"),
    "json_contains": ("expected", "json_contains"),
    "json_paths": ("expected", "json_paths"),
    "schema": ("expected", "schema"),
    "max_duration_ms": ("expected", "max_duration_ms"),
    "duration_ms": ("expected", "max_duration_ms"),
    "response_time_ms": ("expected", "max_duration_ms"),
}


def _select_sheet(workbook: Any, requested: str | None, *, path: Path) -> Any:
    if requested is None:
        return workbook.active
    if requested not in workbook.sheetnames:
        raise CaseDataError(
            f"worksheet does not exist; available={workbook.sheetnames}",
            file=str(path),
            sheet=requested,
        )
    return workbook[requested]


def _normalize_headers(values: Sequence[Any], *, path: Path, sheet: str) -> list[str | None]:
    headers: list[str | None] = []
    seen: set[str] = set()
    for column, value in enumerate(values, start=1):
        if value is None or (isinstance(value, str) and not value.strip()):
            headers.append(None)
            continue
        header = str(value).strip().lower()
        if header in seen:
            raise CaseDataError(
                f"duplicate column header {header!r} at column={column}",
                file=str(path),
                sheet=sheet,
                row=1,
            )
        seen.add(header)
        headers.append(header)
    if not seen:
        raise CaseDataError(
            "worksheet header row is empty",
            file=str(path),
            sheet=sheet,
            row=1,
        )
    return headers


def _raw_row(headers: Sequence[str | None], values: Sequence[Any]) -> dict[str, Any]:
    return {
        header: value
        for header, value in zip(headers, values, strict=False)
        if header is not None and value is not None and value != ""
    }


def _reject_formulas(path: Path, sheet_name: str | None) -> str:
    """Inspect formula expressions before reopening with ``data_only=True``."""

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except (OSError, ValueError) as exc:
        raise CaseDataError(str(exc), file=str(path), sheet=sheet_name) from exc

    try:
        worksheet = _select_sheet(workbook, sheet_name, path=path)
        rows = worksheet.iter_rows()
        first = next(rows, ())
        headers = _normalize_headers(
            [cell.value for cell in first], path=path, sheet=worksheet.title
        )
        id_index = next(
            (index for index, header in enumerate(headers) if header in {"id", "case_id"}),
            None,
        )
        for cells in rows:
            case_id = (
                str(cells[id_index].value).strip()
                if id_index is not None
                and id_index < len(cells)
                and cells[id_index].value is not None
                else None
            )
            for column, cell in enumerate(cells, start=1):
                value = cell.value
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    header = headers[column - 1] if column <= len(headers) else None
                    raise CaseDataError(
                        f"Excel formulas are not allowed (column={header or column})",
                        file=str(path),
                        sheet=worksheet.title,
                        row=cell.row,
                        case_id=case_id,
                    )
        return worksheet.title
    finally:
        workbook.close()


def _decode_json_cell(
    value: Any,
    *,
    field: str,
    path: Path,
    sheet: str,
    row: int,
    case_id: str | None,
) -> Any:
    if not isinstance(value, str):
        return value
    try:
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise CaseDataError(
            f"column {field!r} must contain valid JSON: {exc.msg} (column={exc.colno})",
            file=str(path),
            sheet=sheet,
            row=row,
            case_id=case_id,
        ) from exc


def _decode_markers(
    value: Any,
    *,
    path: Path,
    sheet: str,
    row: int,
    case_id: str | None,
) -> Any:
    if not isinstance(value, str):
        return value
    stripped = value.strip()
    if stripped.startswith("["):
        return _decode_json_cell(
            stripped,
            field="markers",
            path=path,
            sheet=sheet,
            row=row,
            case_id=case_id,
        )
    return [marker.strip() for marker in stripped.split(",") if marker.strip()]


def _set_nested(target: dict[str, Any], section: str, key: str, value: Any) -> None:
    nested = target.setdefault(section, {})
    if not isinstance(nested, dict):
        # This will normally have been caught while parsing the whole section,
        # but keep the failure deterministic if callers pass native cell values.
        raise TypeError(f"{section} must be a JSON object")
    nested[key] = value


def _excel_row_to_mapping(
    raw: Mapping[str, Any], *, path: Path, sheet: str, row: int
) -> dict[str, Any]:
    case_id = case_id_from(raw)
    result: dict[str, Any] = {}

    for key in ("id", "case_id", "name"):
        if key in raw:
            destination = "id" if key == "case_id" else key
            result[destination] = str(raw[key])

    marker_key = "markers" if "markers" in raw else "tags" if "tags" in raw else None
    if marker_key is not None:
        result["markers"] = _decode_markers(
            raw[marker_key],
            path=path,
            sheet=sheet,
            row=row,
            case_id=case_id,
        )

    for section in ("request", "expected"):
        if section not in raw:
            continue
        parsed = _decode_json_cell(
            raw[section],
            field=section,
            path=path,
            sheet=sheet,
            row=row,
            case_id=case_id,
        )
        if not isinstance(parsed, dict):
            raise CaseDataError(
                f"column {section!r} must contain a JSON object",
                file=str(path),
                sheet=sheet,
                row=row,
                case_id=case_id,
            )
        result[section] = parsed

    for field, value in raw.items():
        if field in {"id", "case_id", "name", "markers", "tags", "request", "expected"}:
            continue
        decoded = (
            _decode_json_cell(
                value,
                field=field,
                path=path,
                sheet=sheet,
                row=row,
                case_id=case_id,
            )
            if field in _STRUCTURED_FIELDS
            else value
        )
        if field.startswith("request.") or field.startswith("expected."):
            section, key = field.split(".", 1)
            aliases = {
                "params": "query",
                "body": "json",
                "payload": "json",
                "status_code": "status",
                "json_equals": "json",
                "contains": "json_contains",
                "duration_ms": "max_duration_ms",
                "response_time_ms": "max_duration_ms",
            }
            _set_nested(result, section, aliases.get(key, key), decoded)
        elif field in _FLAT_TARGETS:
            section, key = _FLAT_TARGETS[field]
            _set_nested(result, section, key, decoded)
        else:
            raise CaseDataError(
                f"unsupported column {field!r}",
                file=str(path),
                sheet=sheet,
                row=row,
                case_id=case_id,
            )
    return result


def load_excel_cases(path: str | Path, *, sheet_name: str | None = None) -> list[CaseSpec]:
    """Load cases from an Excel worksheet.

    Formula expressions are rejected in a first pass.  The workbook is then
    reopened with ``data_only=True`` so only stable cell values reach the case
    model.  Both dotted columns (``request.method``) and concise flat columns
    (``method``) are supported.
    """

    source = Path(path)
    actual_sheet = _reject_formulas(source, sheet_name)
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise CaseDataError(str(exc), file=str(source), sheet=actual_sheet) from exc

    try:
        worksheet = _select_sheet(workbook, actual_sheet, path=source)
        values = worksheet.iter_rows(values_only=True)
        header_values = next(values, ())
        headers = _normalize_headers(header_values, path=source, sheet=worksheet.title)
        cases: list[CaseSpec] = []
        seen: dict[str, int] = {}
        for row_number, row_values in enumerate(values, start=2):
            raw = _raw_row(headers, row_values)
            if not raw:
                continue
            mapped = _excel_row_to_mapping(raw, path=source, sheet=worksheet.title, row=row_number)
            case = validate_case_mapping(
                mapped,
                path=source,
                sheet=worksheet.title,
                row=row_number,
            )
            if case.id in seen:
                raise CaseDataError(
                    f"duplicate case id; first declared at row={seen[case.id]}",
                    file=str(source),
                    sheet=worksheet.title,
                    row=row_number,
                    case_id=case.id,
                )
            seen[case.id] = row_number
            cases.append(case)
        return cases
    finally:
        workbook.close()


load_cases = load_excel_cases
