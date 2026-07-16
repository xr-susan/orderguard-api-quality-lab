from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from api_testkit.data import (
    CaseDataError,
    CaseSpec,
    load_excel_cases,
    load_json_cases,
    load_yaml_cases,
)


def _case(case_id: str = "create-order") -> dict[str, object]:
    return {
        "id": case_id,
        "markers": ["smoke", "contract"],
        "request": {
            "method": "post",
            "path": "/orders",
            "headers": {"Idempotency-Key": "case-001"},
            "json": {"sku": "SKU-001", "quantity": 2},
        },
        "expected": {
            "status_code": 201,
            "json_contains": {"status": "CREATED"},
            "json_paths": {"$.quantity": 2},
            "max_duration_ms": 500,
        },
    }


@pytest.mark.parametrize("root_factory", [lambda case: [case], lambda case: {"cases": [case]}])
def test_yaml_loader_accepts_both_supported_root_shapes(
    tmp_path: Path, root_factory: object
) -> None:
    import yaml

    source = tmp_path / "cases.yaml"
    root = root_factory(_case())  # type: ignore[operator]
    source.write_text(yaml.safe_dump(root, allow_unicode=True), encoding="utf-8")

    cases = load_yaml_cases(source)

    assert len(cases) == 1
    case = cases[0]
    assert isinstance(case, CaseSpec)
    assert case.name == "create-order"
    assert case.request.method == "POST"
    assert case.expected.status == 201
    assert case.source is not None
    assert case.source.file == str(source)
    assert case.source.row == 1


def test_json_loader_produces_the_same_canonical_model(tmp_path: Path) -> None:
    source = tmp_path / "cases.json"
    source.write_text(json.dumps({"cases": [_case()]}), encoding="utf-8")

    case = load_json_cases(source)[0]

    assert case.id == "create-order"
    assert case.request.query == {}
    assert case.expected.json_contains == {"status": "CREATED"}


def test_loader_error_identifies_file_row_and_case_id(tmp_path: Path) -> None:
    source = tmp_path / "invalid.json"
    invalid = _case("bad-method")
    invalid["request"] = {"method": "P0ST", "path": "/orders"}
    source.write_text(json.dumps([invalid]), encoding="utf-8")

    with pytest.raises(CaseDataError) as raised:
        load_json_cases(source)

    message = str(raised.value)
    assert str(source) in message
    assert "row=1" in message
    assert "case=bad-method" in message
    assert "request.method" in message


def test_yaml_loader_is_safe_and_rejects_python_tags(tmp_path: Path) -> None:
    source = tmp_path / "unsafe.yaml"
    source.write_text("!!python/object/apply:os.system ['echo unsafe']", encoding="utf-8")

    with pytest.raises(CaseDataError) as raised:
        load_yaml_cases(source)

    assert str(source) in str(raised.value)


def test_duplicate_ids_report_both_current_and_original_rows(tmp_path: Path) -> None:
    source = tmp_path / "duplicate.json"
    source.write_text(json.dumps([_case("same"), _case("same")]), encoding="utf-8")

    with pytest.raises(CaseDataError, match="first declared at row=1") as raised:
        load_json_cases(source)

    assert "row=2" in str(raised.value)
    assert "case=same" in str(raised.value)


def test_excel_loader_supports_flat_and_dotted_columns(tmp_path: Path) -> None:
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "orders"
    worksheet.append(
        [
            "id",
            "name",
            "markers",
            "method",
            "path",
            "request.json",
            "status",
            "expected.json_contains",
            "expected.json_paths",
        ]
    )
    worksheet.append(
        [
            "excel-create",
            "create from workbook",
            "smoke, functional",
            "POST",
            "/orders",
            '{"sku":"SKU-001","quantity":1}',
            201,
            '{"status":"CREATED"}',
            '{"$.quantity":1}',
        ]
    )
    workbook.save(source)
    workbook.close()

    case = load_excel_cases(source, sheet_name="orders")[0]

    assert case.request.json_body == {"sku": "SKU-001", "quantity": 1}
    assert case.expected.status == 201
    assert case.markers == ("smoke", "functional")
    assert case.source is not None
    assert case.source.sheet == "orders"
    assert case.source.row == 2


def test_excel_loader_rejects_formulas_with_exact_source(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "orders"
    worksheet.append(["id", "method", "path", "status"])
    worksheet.append(["formula-case", "GET", "/orders", "=200+0"])
    workbook.save(source)
    workbook.close()

    with pytest.raises(CaseDataError) as raised:
        load_excel_cases(source, sheet_name="orders")

    message = str(raised.value)
    assert str(source) in message
    assert "sheet=orders" in message
    assert "row=2" in message
    assert "case=formula-case" in message
    assert "formulas are not allowed" in message


def test_excel_invalid_json_cell_reports_column_and_case(tmp_path: Path) -> None:
    source = tmp_path / "bad-cell.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "orders"
    worksheet.append(["id", "method", "path", "request.json"])
    worksheet.append(["bad-json", "POST", "/orders", "{broken"])
    workbook.save(source)
    workbook.close()

    with pytest.raises(CaseDataError) as raised:
        load_excel_cases(source)

    message = str(raised.value)
    assert "sheet=orders" in message
    assert "row=2" in message
    assert "case=bad-json" in message
    assert "request.json" in message
